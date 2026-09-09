from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "JanDrishti_SIH26102_SIH_Document.docx"
ORANGE = "C65A32"
INK = "191B1F"
MUTED = "606369"
PALE = "EFE8E0"
TEAL = "1C7369"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False, color=INK, size=10):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(doc, value, level=1):
    p = doc.add_heading(value, level=level)
    p.style.font.name = "Aptos Display" if level == 1 else "Aptos"
    p.style.font.color.rgb = RGBColor.from_string(ORANGE if level == 1 else INK)
    return p


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(value)
    run.font.name = "Aptos"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(INK)


def paragraph(doc, value, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    if bold_prefix and value.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(TEAL)
        r.font.name = "Aptos"
        r.font.size = Pt(10.5)
        value = value[len(bold_prefix):]
    r = p.add_run(value)
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(INK)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("JANDRISHTI")
run.font.name = "Aptos Display"
run.font.size = Pt(36)
run.bold = True
run.font.color.rgb = RGBColor.from_string(INK)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Explainable AI for MPLADS Monitoring, Risk Screening and Accountable Action")
run.font.name = "Aptos"
run.font.size = Pt(16)
run.font.color.rgb = RGBColor.from_string(TEAL)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("Smart India Hackathon 2026  |  Problem Statement SIH26102  |  SIH Solution Document")
run.font.name = "Aptos Mono"
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string(MUTED)

doc.add_paragraph()
callout = doc.add_table(rows=1, cols=1)
callout.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = callout.cell(0, 0)
shade(cell, PALE)
set_cell_text(cell, "Detect  →  Explain  →  Act  →  Learn", bold=True, color=ORANGE, size=16)
cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

heading(doc, "1. Executive Summary")
paragraph(doc, "JanDrishti is an AI-powered monitoring and analytics platform for the Members of Parliament Local Area Development Scheme (MPLADS). It identifies unusual expenditure patterns, duplicate or overlapping works, delayed execution, cost deviations, payment concentration and financial–physical progress mismatches.")
paragraph(doc, "The platform is designed as decision support. It prioritises cases for human review, explains the evidence behind each signal, and connects anomalies to a formal statutory review case, assignment, status workflow and audit trail. It does not automatically declare fraud.")

heading(doc, "2. Problem Statement Alignment")
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for cell, value in zip(table.rows[0].cells, ["SIH requirement", "JanDrishti capability", "Status"]):
    shade(cell, ORANGE)
    set_cell_text(cell, value, True, "FFFFFF", 10)
rows = [
    ("Detect anomalies and irregularities", "Rules, robust statistics, Isolation Forest and payment/network signals", "Implemented"),
    ("Analyse sanctions, expenditure and progress", "Work, transaction, vendor, lifecycle and financial feature layers", "Implemented"),
    ("Find duplicates and cost overruns", "Duplicate candidates, cost deviation and peer benchmarks", "Implemented"),
    ("Generate risk-based alerts", "Severity, 0–100 prioritisation score and evidence", "Implemented"),
    ("Provide predictive insights", "Delay and cost-overrun prioritisation models", "Implemented"),
    ("Support authorities", "Role-based dashboards, Inspect 360° and Cases & Alerts", "Implemented"),
    ("Improve accountability", "Assignments, status changes and audit trail", "Implemented"),
    ("Live detailed official synchronization", "Guarded framework; official granular contract still unavailable", "Partial"),
]
for row in rows:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        set_cell_text(cell, value, color=TEAL if value == "Implemented" else ORANGE, size=9.5)

heading(doc, "3. End-to-End Operating Flow")
for value in [
    "Data ingestion and validation: source checksums, effective dates, batch lineage and quality gates.",
    "AI analysis: deterministic rules, statistical benchmarks, unsupervised ML and network indicators.",
    "Explainable alert: risk score, severity, evidence, thresholds, confidence and limitations.",
    "Inspect 360° dossier: work lifecycle, financial information, progress signals, source provenance and related anomalies.",
    "Statutory review case: assign responsible authority, add directive notes and create a tracked case.",
    "Investigation and action: request clarification, inspect evidence, record status and escalate where required.",
    "Closure and learning: record resolution evidence, reviewer decision and expert label for future model evaluation.",
]:
    bullet(doc, value)

heading(doc, "4. AI and Analytics")
ai_table = doc.add_table(rows=1, cols=2)
ai_table.style = "Table Grid"
for cell, value in zip(ai_table.rows[0].cells, ["Signal layer", "Examples"]):
    shade(cell, TEAL)
    set_cell_text(cell, value, True, "FFFFFF", 10)
for row in [
    ("Deterministic", "Expenditure limits, overdue duration, missing information and compliance checks."),
    ("Statistical", "Percentiles, robust Z-scores, MAD outliers and category/state peer comparisons."),
    ("Unsupervised ML", "Isolation Forest for unusual multi-dimensional work and transaction patterns."),
    ("Duplicate detection", "Text similarity, fuzzy matching, cost proximity and geographic candidate review."),
    ("Payment/network", "Payment velocity, vendor concentration, repeated amounts and year-end clustering."),
    ("Predictive", "Chronological delay and cost-overrun models with uncertainty estimates."),
]:
    cells = ai_table.add_row().cells
    set_cell_text(cells[0], row[0], True, ORANGE, 9.5)
    set_cell_text(cells[1], row[1], size=9.5)
paragraph(doc, "Governance principle: the platform reports potential anomaly or risk-prioritisation signals. Confirmed fraud or statutory findings require authorised human investigation.")

heading(doc, "5. Role-Based Workflow")
role_table = doc.add_table(rows=1, cols=3)
role_table.style = "Table Grid"
for cell, value in zip(role_table.rows[0].cells, ["Role", "Primary responsibility", "Current implementation"]):
    shade(cell, ORANGE)
    set_cell_text(cell, value, True, "FFFFFF", 10)
for row in [
    ("Member of Parliament", "View constituency works, risk and case status; request follow-up.", "Implemented for scoped viewing."),
    ("District Authority", "Triage local alerts, initiate cases, assign officers and update status.", "Implemented."),
    ("State Nodal Authority", "Compare districts, monitor patterns and coordinate escalation.", "Implemented for monitoring; automated escalation is future work."),
    ("Ministry / MoSPI", "Review national trends, high-risk cases and policy signals.", "Implemented for oversight dashboards."),
    ("Auditor / Analyst", "Examine evidence, challenge signals and provide expert labels.", "Implemented; label collection must continue."),
    ("Citizen", "View public information and submit observations/evidence.", "Implemented."),
    ("Implementing Agency", "Respond to clarifications and provide official evidence.", "Dedicated workflow remains future work."),
]:
    cells = role_table.add_row().cells
    for cell, value in zip(cells, row):
        set_cell_text(cell, value, size=9.2, color=TEAL if value.startswith("Implemented") else ORANGE)

heading(doc, "6. Current Product Workflow")
paragraph(doc, "A typical demonstration begins with a District Authority login. The user opens the AI Anomaly Center, selects a high-risk work, opens the Inspect 360° dossier, reviews the score breakdown and evidence, then selects Initiate Statutory Review Case. The case is assigned, appears in Cases & Alerts, and every status change is written to the audit trail.")
for value in [
    "AI anomaly detected",
    "Inspect 360° opened",
    "Severity and assignee selected",
    "Statutory review case created",
    "Case status and notes updated",
    "Audit trail reviewed",
]:
    bullet(doc, value)

heading(doc, "7. Data Trust and Provenance")
for value in [
    "Required-input preflight checks prevent incomplete database builds.",
    "Data-quality validation checks IDs, amounts, durations and orphan references.",
    "Source checksum, effective date, batch ID and source health are retained.",
    "Aggregate MoSPI telemetry synchronization is live.",
    "Detailed works, payments, vendors and derived anomalies remain snapshot-backed because no usable official granular API/export contract is currently available.",
]:
    bullet(doc, value)

heading(doc, "8. Implementation Status")
status_table = doc.add_table(rows=1, cols=2)
status_table.style = "Table Grid"
for cell, value in zip(status_table.rows[0].cells, ["Area", "Status"]):
    shade(cell, TEAL)
    set_cell_text(cell, value, True, "FFFFFF", 10)
for row in [
    ("Explainable anomaly detection", "Implemented"),
    ("AI feature snapshots and governance", "Implemented"),
    ("Drift monitoring and deterministic fixtures", "Implemented"),
    ("Delay and cost-overrun predictive insights", "Implemented"),
    ("Human review labels and guarded fraud model", "Implemented; training gated until labels exist"),
    ("Inspect 360° and statutory review case creation", "Implemented"),
    ("Cases, alerts, assignments and audit trail", "Implemented"),
    ("Live detailed official data", "Blocked by missing official granular contract"),
]:
    cells = status_table.add_row().cells
    set_cell_text(cells[0], row[0], size=9.5)
    set_cell_text(cells[1], row[1], True, ORANGE if "Blocked" in row[1] else TEAL, 9.5)

heading(doc, "9. Judge Demonstration Script")
for value in [
    "Login as District Authority.",
    "Open AI Anomaly Center and choose a high-risk signal.",
    "Open Inspect 360° and explain the risk score and evidence.",
    "Initiate Statutory Review Case with assignee and directive notes.",
    "Open Cases & Alerts, update the case status and show the audit entry.",
    "Switch to Citizen view and demonstrate transparency and public reporting.",
]:
    bullet(doc, value)

heading(doc, "10. Roadmap")
for value in [
    "Add official granular synchronization when an authenticated, documented portal API/export is provided.",
    "Add case evidence uploads with document type, uploader, timestamp, hash and verification status.",
    "Add inspection scheduling, deadlines, notifications and automatic escalation.",
    "Require resolution evidence and authority approval before closure.",
    "Collect qualified review labels and then evaluate supervised fraud-risk prioritisation.",
]:
    bullet(doc, value)

heading(doc, "11. Closing Statement")
paragraph(doc, "JanDrishti addresses the SIH26102 requirement by connecting AI-powered anomaly detection to explainable decision support and an accountable administrative workflow. Its central value is not simply finding unusual records; it helps the right authority understand the signal, investigate it, take action and record what happened.")

doc.save(OUT)
print(OUT)
