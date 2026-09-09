from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


OUT = "JanDrishti_SIH26102_Judge_Deck.pptx"
BG = RGBColor(248, 247, 243)
INK = RGBColor(25, 27, 31)
MUTED = RGBColor(96, 99, 105)
ORANGE = RGBColor(198, 90, 50)
TEAL = RGBColor(28, 115, 105)
PALE = RGBColor(239, 232, 224)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def box(slide, x, y, w, h, fill=WHITE, line=PALE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.7)
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font="Aptos",
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def bullet_list(slide, items, x, y, w, h, size=16, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(10)
        p.bullet = True
    return shape


def header(slide, kicker, title_value, subtitle=None):
    text(slide, kicker.upper(), 0.65, 0.38, 5.8, 0.25, 10, ORANGE, True, "Aptos Mono")
    text(slide, title_value, 0.65, 0.78, 12.0, 0.62, 29, INK, True, "Georgia")
    if subtitle:
        text(slide, subtitle, 0.68, 1.48, 11.8, 0.42, 12, MUTED)


def footer(slide, number):
    text(slide, f"JAN DRISHTI  •  SIH26102                                      {number:02d}",
         0.68, 7.12, 12, 0.18, 8, MUTED, False, "Aptos Mono")


def add_slide(kicker, title_value, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    header(slide, kicker, title_value, subtitle)
    footer(slide, len(prs.slides))
    return slide


def card(slide, title_value, body, x, y, w, h, accent=ORANGE):
    box(slide, x, y, w, h)
    box(slide, x, y, 0.08, h, accent, accent, False)
    text(slide, title_value, x + 0.22, y + 0.18, w - 0.4, 0.32, 15, INK, True)
    text(slide, body, x + 0.22, y + 0.63, w - 0.42, h - 0.78, 12, MUTED)


# 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG
text(slide, "SMART INDIA HACKATHON 2026  /  PROBLEM STATEMENT SIH26102", 0.75, 0.65, 8, 0.25, 11, ORANGE, True, "Aptos Mono")
text(slide, "JanDrishti", 0.72, 1.35, 8.7, 1.1, 54, INK, True, "Georgia")
text(slide, "Explainable AI for MPLADS monitoring,\nrisk screening and accountable action.", 0.78, 2.55, 8.5, 1.0, 25, MUTED, False, "Georgia")
box(slide, 9.65, 1.25, 2.8, 3.7, ORANGE, ORANGE)
text(slide, "DETECT", 10.05, 1.75, 2.0, 0.35, 18, WHITE, True, "Aptos Mono", PP_ALIGN.CENTER)
text(slide, "EXPLAIN", 10.05, 2.55, 2.0, 0.35, 18, WHITE, True, "Aptos Mono", PP_ALIGN.CENTER)
text(slide, "ACT", 10.05, 3.35, 2.0, 0.35, 18, WHITE, True, "Aptos Mono", PP_ALIGN.CENTER)
text(slide, "LEARN", 10.05, 4.15, 2.0, 0.35, 18, WHITE, True, "Aptos Mono", PP_ALIGN.CENTER)
text(slide, "MPLADS • transparency • early warning • human review", 0.78, 6.25, 8.5, 0.3, 13, TEAL, True)
footer(slide, 1)

# 2
slide = add_slide("01  /  THE PROBLEM", "MPLADS data is large, fragmented and difficult to supervise",
                  "Thousands of works, payments and implementing agencies create a monitoring problem—not only a reporting problem.")
card(slide, "Scale", "Works, sanctions, payments, vendors and progress records span districts and constituencies.", 0.7, 2.15, 3.8, 2.0)
card(slide, "Risk", "Duplicate proposals, unusual costs, delayed execution and payment–progress divergence can remain hidden.", 4.75, 2.15, 3.8, 2.0, TEAL)
card(slide, "Gap", "Authorities need prioritised, explainable cases—not another static spreadsheet.", 8.8, 2.15, 3.8, 2.0, ORANGE)
text(slide, "Core question: Which work needs attention first, why, and what should happen next?", 1.0, 5.2, 11.2, 0.55, 23, INK, True, "Georgia", PP_ALIGN.CENTER)

# 3
slide = add_slide("02  /  OUR SOLUTION", "One monitoring loop from signal to accountable action",
                  "JanDrishti converts analytical signals into a governed review workflow.")
steps = [("01", "Ingest", "Official observations + validated datasets"),
         ("02", "Analyse", "Rules, statistics, ML and graph signals"),
         ("03", "Explain", "Evidence, thresholds, confidence and limitations"),
         ("04", "Act", "Inspect 360° → initiate review case"),
         ("05", "Learn", "Reviewer outcome becomes labelled feedback")]
for i, (num, title_value, body) in enumerate(steps):
    x = 0.7 + i * 2.5
    box(slide, x, 2.3, 2.15, 2.5, WHITE, PALE)
    text(slide, num, x + 0.2, 2.55, 0.5, 0.35, 18, ORANGE, True, "Aptos Mono")
    text(slide, title_value, x + 0.2, 3.05, 1.75, 0.35, 16, INK, True)
    text(slide, body, x + 0.2, 3.6, 1.75, 0.8, 11, MUTED)

# 4
slide = add_slide("03  /  PLATFORM", "Decision-support architecture built for public accountability")
layers = [("Experience", "React + TypeScript dashboards\nRole-based workspaces • Inspect 360° • Cases & Alerts"),
          ("Intelligence", "Explainable anomaly ensemble\nDelay and cost-overrun predictions\nDuplicate and payment-network signals"),
          ("Governance", "Review cases • audit trail • human labels\nJurisdiction controls • evidence workflow"),
          ("Data", "FastAPI services • SQLite/PostgreSQL compatibility\nSource provenance • validation gates • batch lineage")]
for i, (title_value, body) in enumerate(layers):
    y = 2.05 + i * 1.05
    box(slide, 1.0, y, 11.3, 0.78, WHITE if i % 2 == 0 else PALE, PALE)
    text(slide, title_value, 1.3, y + 0.18, 2.0, 0.28, 15, ORANGE if i == 1 else TEAL, True)
    text(slide, body, 3.4, y + 0.13, 8.2, 0.42, 13, INK)

# 5
slide = add_slide("04  /  AI", "AI detects risk signals—never declares guilt")
card(slide, "Deterministic", "Cost variance, overdue duration, expenditure limits, missing information and compliance checks.", 0.75, 2.1, 3.0, 2.35, ORANGE)
card(slide, "Statistical", "Robust benchmarks, percentiles, MAD/Z-score outliers and peer comparisons.", 3.95, 2.1, 3.0, 2.35, TEAL)
card(slide, "Unsupervised ML", "Isolation Forest identifies unusual multi-dimensional work and transaction patterns.", 7.15, 2.1, 3.0, 2.35, ORANGE)
card(slide, "Network signals", "Payment velocity, vendor concentration, repeated amounts and financial-year-end clustering.", 10.35, 2.1, 2.3, 2.35, TEAL)
box(slide, 1.0, 5.1, 11.2, 0.75, PALE, PALE)
text(slide, "Output: a 0–100 prioritisation score + evidence + limitations + recommended human review.", 1.25, 5.32, 10.7, 0.3, 17, INK, True, "Georgia", PP_ALIGN.CENTER)

# 6
slide = add_slide("05  /  ROLE WORKFLOW", "Every role sees the action appropriate to its mandate")
roles = [("MP", "View constituency risk\nrequest follow-up"),
         ("District", "Triage, assign, inspect\nupdate review case"),
         ("State", "Compare districts\nescalate patterns"),
         ("Ministry", "National trends\npolicy and oversight"),
         ("Auditor", "Independent evidence\nreview and labels"),
         ("Citizen", "Public view\nsubmit observations")]
for i, (title_value, body) in enumerate(roles):
    x = 0.7 + (i % 3) * 4.1
    y = 2.05 + (i // 3) * 2.0
    card(slide, title_value, body, x, y, 3.65, 1.45, TEAL if i % 2 else ORANGE)
text(slide, "AI anomaly → Inspect 360° → Initiate Statutory Review Case → Cases & Alerts → Audit trail → Resolution", 0.9, 6.25, 11.5, 0.35, 16, INK, True, "Georgia", PP_ALIGN.CENTER)

# 7
slide = add_slide("06  /  THE CORE DEMO", "Inspect 360° turns an alert into an accountable case")
demo_steps = [("1", "Open anomaly", "Review score, reasons and affected work"),
              ("2", "Inspect dossier", "See lifecycle, finance, progress and evidence"),
              ("3", "Initiate case", "Set severity, assignee and directive notes"),
              ("4", "Track action", "Status changes and immutable audit entries")]
for i, (n, t, b) in enumerate(demo_steps):
    x = 0.8 + i * 3.1
    box(slide, x, 2.25, 2.65, 2.55, WHITE, PALE)
    text(slide, n, x + 0.22, 2.52, 0.4, 0.4, 24, ORANGE, True, "Georgia")
    text(slide, t, x + 0.22, 3.2, 2.1, 0.3, 16, INK, True)
    text(slide, b, x + 0.22, 3.75, 2.1, 0.7, 12, MUTED)
text(slide, "No black-box alert disappears into a dashboard: it becomes a traceable administrative decision.", 1.0, 5.7, 11.2, 0.4, 20, TEAL, True, "Georgia", PP_ALIGN.CENTER)

# 8
slide = add_slide("07  /  DATA TRUST", "Safe publication and honest provenance")
card(slide, "Validation gates", "Required inputs, unique IDs, negative amounts, invalid durations and orphan references are checked before publication.", 0.75, 2.15, 3.8, 2.35, ORANGE)
card(slide, "Lineage", "Source checksum, effective date, batch ID, source health and archived observations are retained.", 4.78, 2.15, 3.8, 2.35, TEAL)
card(slide, "Honest status", "Aggregate MoSPI telemetry is live. Detailed works, payments, vendors and derived anomalies remain snapshot-backed until an official granular contract is available.", 8.8, 2.15, 3.8, 2.35, ORANGE)
text(slide, "This distinction protects decision-makers from false confidence.", 1.0, 5.35, 11.2, 0.45, 22, INK, True, "Georgia", PP_ALIGN.CENTER)

# 9
slide = add_slide("08  /  IMPLEMENTATION", "What is working today")
items = ["Explainable anomaly center with duplicate, mismatch, delay and outlier views",
         "Inspect 360° dossier with statutory review case initiation",
         "Cases & Alerts with assignment, status transitions and audit trail",
         "AI feature snapshots, model governance and drift monitoring",
         "Delay and cost-overrun predictive insights with leakage controls",
         "Expert-label workflow guarded against unsupported fraud claims"]
bullet_list(slide, items, 0.95, 2.0, 11.3, 3.8, 17)
box(slide, 1.0, 5.95, 11.2, 0.62, PALE, PALE)
text(slide, "Validated locally: frontend production build passes • data-quality validation passes • AI fixture checks 4/4", 1.25, 6.14, 10.7, 0.25, 13, INK, True, "Aptos Mono", PP_ALIGN.CENTER)

# 10
slide = add_slide("09  /  IMPACT", "From manual monitoring to risk-based attention")
impact = [("Earlier warning", "Surface unusual patterns before routine review catches them."),
          ("Less noise", "Prioritise cases by severity, evidence and jurisdiction."),
          ("Better accountability", "Every action has an owner, timestamp and audit record."),
          ("Public confidence", "Citizens can see project information and submit observations.")]
for i, (t, b) in enumerate(impact):
    y = 2.0 + i * 0.95
    text(slide, f"0{i+1}", 1.0, y, 0.6, 0.3, 16, ORANGE, True, "Aptos Mono")
    text(slide, t, 1.8, y, 2.5, 0.3, 17, INK, True)
    text(slide, b, 4.55, y, 7.4, 0.35, 15, MUTED)

# 11
slide = add_slide("10  /  DEMO SCRIPT", "A five-minute judge walkthrough")
demo = ["Login as District Authority",
        "Open AI Anomaly Center and select a high-risk work",
        "Open Inspect 360° and review evidence and score breakdown",
        "Initiate Statutory Review Case and assign the district authority",
        "Update status to inspection / clarification and show audit trail",
        "Switch to Citizen view and show public transparency + reporting"]
bullet_list(slide, demo, 1.0, 2.0, 10.8, 4.3, 18)

# 12
slide = add_slide("11  /  ROADMAP", "The next layer is operational depth—not another dashboard")
card(slide, "Now", "Role-based review, explainable risk, provenance, cases and audit trail.", 0.8, 2.0, 3.7, 2.0, TEAL)
card(slide, "Next", "Case evidence uploads, inspection scheduling, deadlines, notifications and closure validation.", 4.8, 2.0, 3.7, 2.0, ORANGE)
card(slide, "Scale", "Official granular synchronization, more reviewed labels and calibrated supervised prioritisation.", 8.8, 2.0, 3.7, 2.0, TEAL)
text(slide, "JanDrishti is designed to help authorities act earlier, explain decisions clearly and learn from every reviewed case.", 1.0, 5.2, 11.2, 0.7, 23, INK, True, "Georgia", PP_ALIGN.CENTER)
text(slide, "Thank you  •  Questions", 4.4, 6.35, 4.5, 0.3, 16, ORANGE, True, "Aptos Mono", PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
